#!/usr/bin/env python3
"""Export a validated quiz (Emile JSON contract) to Canvas QTI 1.2 and/or Word.

This is a human-run step, deliberately outside the MCP server so the server
stays read-only. It refuses to export a quiz that fails `validate_quiz`
unless you pass --force.

    python scripts/export_quiz.py quiz.json                 # -> quiz.qti.zip and quiz.docx
    python scripts/export_quiz.py quiz.json --format qti    # QTI only (stdlib, no extra deps)
    python scripts/export_quiz.py quiz.json --format docx   # needs: pip install -r requirements-export.txt

Import the .zip in Canvas via Settings -> Import Course Content -> QTI .zip file.
"""

import argparse
import html
import json
import re
import sys
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
import validators  # noqa: E402

CANVAS_TYPES = {
    "true_false": "true_false_question",
    "multiple_choice": "multiple_choice_question",
    "multiple_answer": "multiple_answers_question",
    "fill_blank": "short_answer_question",
    "manual": "essay_question",
}
TF_OPTIONS = ["True", "False"]
# Spec issues --force may override; anything else means the item can't be rendered.
OVERRIDABLE = re.compile(
    rf"^(total|{'|'.join(validators.TYPES)}|{'|'.join(validators.LEVELS)}):"
    r"|^Q\d+: (no objective mapped|all/none of the above|possible double negative|unknown level)"
)


def _sub(parent, tag, text=None, **attrs):
    el = ET.SubElement(parent, tag, {k.rstrip("_"): str(v) for k, v in attrs.items()})
    if text is not None:
        el.text = text
    return el


def _material(parent, text, html_text=True):
    mat = _sub(parent, "material")
    _sub(mat, "mattext", f"<p>{html.escape(text)}</p>" if html_text else text,
         texttype="text/html" if html_text else "text/plain")


def _meta(item, fields):
    qm = _sub(_sub(item, "itemmetadata"), "qtimetadata")
    for label, entry in fields.items():
        f = _sub(qm, "qtimetadatafield")
        _sub(f, "fieldlabel", label)
        _sub(f, "fieldentry", str(entry))


def _item(section, q, n):
    qtype = q["type"]
    ident = f"q{n}"
    item = _sub(section, "item", ident=ident, title=f"Question {n}")
    _meta(item, {"question_type": CANVAS_TYPES[qtype], "points_possible": "1",
                 "objective": q.get("objective", ""), "level": q.get("level", "")})
    pres = _sub(item, "presentation")
    _material(pres, q["stem"])

    res = _sub(item, "resprocessing")
    _sub(_sub(res, "outcomes"), "decvar", maxvalue="100", minvalue="0", varname="SCORE", vartype="Decimal")
    if q.get("rationale"):
        cond = _sub(res, "respcondition", continue_="Yes")
        _sub(_sub(cond, "conditionvar"), "other")
        _sub(cond, "displayfeedback", feedbacktype="Response", linkrefid="general_fb")

    if qtype in ("true_false", "multiple_choice", "multiple_answer"):
        opts = q.get("options") or TF_OPTIONS
        ids = [f"{ident}_a{i}" for i in range(len(opts))]
        lid = _sub(pres, "response_lid", ident="response1",
                   rcardinality="Multiple" if qtype == "multiple_answer" else "Single")
        choice = _sub(lid, "render_choice")
        for oid, text in zip(ids, opts):
            _material(_sub(choice, "response_label", ident=oid), text, html_text=False)
        cond = _sub(res, "respcondition", continue_="No")
        var = _sub(cond, "conditionvar")
        if qtype == "multiple_answer":
            both = _sub(var, "and")
            for i, oid in enumerate(ids):
                target = both if i in q["answer"] else _sub(both, "not")
                _sub(target, "varequal", oid, respident="response1")
        else:
            _sub(var, "varequal", ids[q["answer"][0]], respident="response1")
        _sub(cond, "setvar", "100", action="Set", varname="SCORE")
    elif qtype == "fill_blank":
        rs = _sub(pres, "response_str", ident="response1", rcardinality="Single")
        _sub(_sub(rs, "render_fib"), "response_label", ident="answer1", rshuffle="No")
        cond = _sub(res, "respcondition", continue_="No")
        var = _sub(cond, "conditionvar")
        for accepted in q["answer"]:
            _sub(var, "varequal", accepted, respident="response1")
        _sub(cond, "setvar", "100", action="Set", varname="SCORE")
    else:  # manual
        rs = _sub(pres, "response_str", ident="response1", rcardinality="Single")
        _sub(_sub(rs, "render_fib"), "response_label", ident="answer1", rshuffle="No")
        _sub(_sub(_sub(res, "respcondition", continue_="No"), "conditionvar"), "other")

    if q.get("rationale"):
        fb = _sub(item, "itemfeedback", ident="general_fb")
        _material(_sub(fb, "flow_mat"), q["rationale"])


def build_qti(quiz: dict, out: Path) -> Path:
    aid = "g" + uuid.uuid4().hex
    root = ET.Element("questestinterop", xmlns="http://www.imsglobal.org/xsd/ims_qtiasiv1p2")
    assessment = _sub(root, "assessment", ident=aid, title=quiz.get("title", "Quiz"))
    section = _sub(assessment, "section", ident="root_section")
    for n, q in enumerate(quiz["questions"], 1):
        _item(section, q, n)

    manifest = ET.Element("manifest", identifier="m" + aid,
                          xmlns="http://www.imsglobal.org/xsd/imsccv1p1/imscp_v1p1")
    md = _sub(manifest, "metadata")
    _sub(md, "schema", "IMS Content")
    _sub(md, "schemaversion", "1.1.3")
    _sub(manifest, "organizations")
    resource = _sub(_sub(manifest, "resources"), "resource", identifier=aid, type="imsqti_xmlv1p2")
    _sub(resource, "file", href=f"{aid}/{aid}.xml")

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("imsmanifest.xml", ET.tostring(manifest, encoding="unicode", xml_declaration=True))
        z.writestr(f"{aid}/{aid}.xml", ET.tostring(root, encoding="unicode", xml_declaration=True))
    return out


def build_docx(quiz: dict, out: Path) -> Path:
    try:
        from docx import Document
    except ImportError:
        sys.exit("Word export needs python-docx: pip install -r requirements-export.txt")

    doc = Document()
    doc.add_heading(quiz.get("title", "Quiz"), level=1)
    if quiz.get("course"):
        doc.add_paragraph(quiz["course"])
    for n, q in enumerate(quiz["questions"], 1):
        doc.add_paragraph(f"{n}. {q['stem']}")
        opts = q.get("options") or (TF_OPTIONS if q["type"] == "true_false" else [])
        for i, text in enumerate(opts):
            doc.add_paragraph(f"{chr(65 + i)}. {text}", style="List Paragraph")
        if q["type"] == "fill_blank":
            doc.add_paragraph("Answer: ____________________")
        if q["type"] == "manual":
            doc.add_paragraph("\n\n")

    doc.add_page_break()
    doc.add_heading("Answer Key", level=1)
    for n, q in enumerate(quiz["questions"], 1):
        if q["type"] == "manual":
            key = "Manually graded"
        elif q["type"] == "fill_blank":
            key = " / ".join(q["answer"])
        else:
            key = ", ".join(chr(65 + a) for a in q["answer"])
        doc.add_paragraph(f"{n}. {key}  [{q.get('level', '')}; {q.get('objective', '')}]")
        if q.get("rationale"):
            doc.add_paragraph(q["rationale"], style="Intense Quote")
    doc.save(out)
    return out


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("quiz", type=Path, help="quiz JSON file (Emile contract)")
    ap.add_argument("--format", choices=["qti", "docx", "both"], default="both")
    ap.add_argument("--out-dir", type=Path, help="default: next to the input file")
    ap.add_argument("--force", action="store_true", help="export even if validation fails")
    args = ap.parse_args()

    raw = args.quiz.read_text(encoding="utf-8")
    result = validators.validate_quiz(raw)
    if not result["pass"]:
        print("Validation issues:", *(f"  - {i}" for i in result["issues"]), sep="\n", file=sys.stderr)
        if not args.force:
            sys.exit("Not exported. Fix the issues or pass --force.")
        # --force overrides spec rules, but the exporter still needs items it can render.
        if any(not OVERRIDABLE.match(i) for i in result["issues"]):
            sys.exit("Not exported: some items are malformed (type, stem, options, or answer), even with --force.")

    quiz = json.loads(raw)
    out_dir = args.out_dir or args.quiz.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = args.quiz.stem
    if args.format in ("qti", "both"):
        print("Wrote", build_qti(quiz, out_dir / f"{stem}.qti.zip"))
    if args.format in ("docx", "both"):
        print("Wrote", build_docx(quiz, out_dir / f"{stem}.docx"))


if __name__ == "__main__":
    main()
